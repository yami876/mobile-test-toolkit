# -*- coding: utf-8 -*-
"""
测试机信息采集脚本（Android + iOS，列格式与测试机资产表一致）
iOS 采集见 ios采集.py，需 pip install pymobiledevice3

用法：
    1. 手机打开「开发者选项 → USB 调试」，用数据线连电脑，手机上点「允许」
    2. python 采集设备信息.py            # 采集当前所有已连接的设备
       python 采集设备信息.py -o xx.xlsx  # 写入指定的 Excel（表头需与下方「表头」一致）
    3. 结果写入同目录的「测试机信息.xlsx」，按序列号去重，重复采集会更新旧行

说明：
    - 前 17 列与现有资产表一致（原「型号/序列号」拆为两列），后面追加 App 测试维度和借用管理列
    - 手工维护列（见「手工列」）脚本不填、也不覆盖
    - 主屏尺寸、蓝牙版本系统无法可靠读取，按型号查官方参数表填写
    - 运行内存、存储容量按读数向上取整为标称值（如读到 7.4G → 8G）
    - 系统mac地址取 WLAN 出厂地址；读不到时取当前地址，可能为随机 MAC，需人工核对
"""
import argparse
import math
import os
import re
import subprocess
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# 输出被重定向到 GBK 终端/文件时，遇到无法编码的符号不报错
sys.stdout.reconfigure(errors="replace")

脚本目录 = os.path.dirname(os.path.abspath(__file__))
默认输出 = os.path.join(脚本目录, "测试机信息.xlsx")

# 前 17 列与现有资产表一致（原「型号/序列号」拆为两列）
现有列 = [
    "资产编号", "设备型号", "资产所有人", "设备类型", "型号代号", "序列号", "操作系统",
    "CPU", "GPU", "运行内存", "存储容量", "主屏尺寸", "分辨率", "蓝牙",
    "系统mac地址", "其他参数", "上市日期",
]
# 新增列：App 测试维度 + 借用管理
新增列 = [
    "API级别", "安全补丁", "CPU架构", "谷歌服务", "WebView版本", "刷新率", "屏幕密度",
    "蓝牙MAC", "UDID", "版本区域", "网络版本", "性能档位", "屏幕形态", "手写笔", "随附配件", "登录账号开机密码", "当前借用人", "设备状态", "采集备注",
]
表头 = 现有列 + 新增列
列宽 = [12, 18, 10, 9, 14, 20, 24, 26, 18, 9, 9, 9, 12, 8, 19, 20, 11,
        8, 11, 9, 9, 16, 8, 9, 19, 28, 9, 9, 9, 11, 16, 16, 20, 11, 9, 30]

# 手工维护列：脚本不填，已有内容时也不覆盖
手工列 = {"资产编号", "资产所有人", "其他参数", "随附配件", "登录账号开机密码", "当前借用人", "设备状态"}
# 建议列：脚本给出初值，仅在格子为空时填写，人工改过的不覆盖
# （主屏尺寸/蓝牙/上市日期 Android 读不到需手工填，iOS 由内置机型表给出）
建议列 = {"版本区域", "网络版本", "性能档位", "屏幕形态", "主屏尺寸", "蓝牙", "上市日期", "手写笔"}
下拉选项 = {
    "版本区域": ["国行", "海外"],
    "网络版本": ["WLAN版", "蜂窝版"],
    "性能档位": ["高端", "中端", "低端"],
    "屏幕形态": ["挖孔", "水滴", "刘海", "灵动岛", "Home键", "全面屏无开孔", "折叠屏", "平板"],
    "设备状态": ["在库", "借出", "维修", "报废"],
    "手写笔": ["支持（随机附带）", "支持（需另购）", "支持", "不支持"],
}

# 芯片代号 → 市场名称（未收录的保留原代号，标注「待查」）
芯片名称表 = {
    # 高通 SoC 型号
    "SM8850": "第五代骁龙8至尊版", "SM8750": "骁龙8至尊版", "SM8650": "骁龙8 Gen3", "SM8635": "骁龙8s Gen3",
    "SM8550": "骁龙8 Gen2", "SM8475": "骁龙8+ Gen1", "SM8450": "骁龙8 Gen1",
    "SM8350": "骁龙888", "SM8250": "骁龙865/870", "SM8150": "骁龙855",
    "SDM845": "骁龙845", "SM7675": "骁龙7+ Gen3", "SM7550": "骁龙7 Gen3",
    "SM7475": "骁龙7+ Gen2", "SM7325": "骁龙778G", "SM7225": "骁龙750G", "SM6375": "骁龙695",
    "SM6225": "骁龙680",
    # 高通平台代号
    "sun": "骁龙8至尊版", "pineapple": "骁龙8 Gen3", "kalama": "骁龙8 Gen2",
    "taro": "骁龙8 Gen1/8+ Gen1", "lahaina": "骁龙888/778G", "kona": "骁龙865/870",
    "msmnile": "骁龙855", "sdm660": "骁龙660", "SDM660": "骁龙660",
    # 联发科
    "MT6991": "天玑9400/9400+", "MT6989": "天玑9300", "MT6985": "天玑9200",
    "MT6983": "天玑9000", "MT6896": "天玑8200", "MT6895": "天玑8100",
    "MT6893": "天玑1200", "MT6789": "Helio G99",
    "MT6899": "天玑8500", "MT6897": "天玑8300", "MT8792": "天玑8300", "MT8775": "天玑7300",
    "MT8755": "天玑6300", "MT6835": "天玑6100+/6300",  # MT87xx 为平板版代号
    # 谷歌 / 三星
    "gs101": "Google Tensor", "gs201": "Google Tensor G2",
    "zuma": "Google Tensor G3", "zumapro": "Google Tensor G4",
    "mt8169": "MediaTek MT8169",
    # 华为海思（平台代号）
    "orlando": "麒麟810", "hi3660": "麒麟960",
    "s5e9925": "Exynos 2200", "s5e8855": "Exynos 1580", "Exynos 9611": "Exynos 9611",
}

# 芯片市场名称 → 性能档位（未收录的留空，人工判断）
芯片档位表 = {
    "高端": ["第五代骁龙8至尊版", "骁龙8至尊版", "骁龙8 Gen3", "骁龙8s Gen3", "骁龙8 Gen2", "骁龙8+ Gen1",
             "骁龙8 Gen1", "骁龙8 Gen1/8+ Gen1", "天玑9400", "天玑9400/9400+", "天玑9300", "天玑9200",
             "天玑9000", "Google Tensor G2", "Google Tensor G3", "Google Tensor G4"],
    "中端": ["骁龙888", "骁龙888/778G", "骁龙888+", "麒麟810", "骁龙750G", "骁龙865/870", "骁龙870", "骁龙865+", "骁龙865", "骁龙855", "骁龙845", "骁龙7+ Gen3",
             "骁龙7 Gen3", "骁龙7+ Gen2", "骁龙778G", "天玑7300", "天玑8500", "天玑8500-Ultra", "天玑8300", "天玑8200", "天玑8100", "天玑1200",
             "Google Tensor", "Exynos 2200", "Exynos 1580"],
    "低端": ["骁龙695", "骁龙685", "麒麟960", "骁龙680", "骁龙660", "Helio G99", "MediaTek MT8169", "Exynos 9611", "天玑6300", "天玑6100+/6300"],
}

# 亚马逊 Fire 平板型号代号 → 市场名
亚马逊型号表 = {"KFTUWI": "Fire HD 10（第13代）"}

# 不设 10G 档：三星等预留内存较多，12G 机型系统内只看得到不到 10G，会被误判成 10G
内存档位 = [1, 2, 3, 4, 6, 8, 12, 16, 18, 20, 24, 32]
存储档位 = [8, 16, 32, 64, 128, 256, 512, 1024, 2048]
MAC正则 = r"(([0-9A-Fa-f]{2}:){5}[0-9A-Fa-f]{2})"


def adb(序列号, *命令, 超时=20):
    """执行 adb 命令，失败返回空字符串"""
    参数 = ["adb"] + (["-s", 序列号] if 序列号 else []) + list(命令)
    try:
        结果 = subprocess.run(参数, capture_output=True, timeout=超时)
        return 结果.stdout.decode("utf-8", errors="replace").strip()
    except Exception:
        return ""


def shell(序列号, 命令):
    return adb(序列号, "shell", 命令)


def 取属性(序列号, *属性名):
    """按顺序取第一个非空的系统属性"""
    for 名 in 属性名:
        值 = shell(序列号, f"getprop {名}")
        if 值:
            return 值
    return ""


def 向上取档(数值, 档位):
    for 档 in 档位:
        if 数值 <= 档:
            return 档
    return math.ceil(数值)


def 格式化容量(GB):
    return f"{GB // 1024}T" if GB >= 1024 else f"{GB}G"


def 是有效MAC(值):
    return bool(re.fullmatch(MAC正则, 值 or "")) \
        and 值.upper() not in ("02:00:00:00:00:00", "00:00:00:00:00:00")


def 采集CPU(序列号):
    厂商 = 取属性(序列号, "ro.soc.manufacturer")
    型号 = 取属性(序列号, "ro.soc.model")
    平台 = 取属性(序列号, "ro.board.platform", "ro.hardware")
    # 骁龙865/865+/870 共用 SM8250（kona），按大核最高频率区分
    if 型号.upper() == "SM8250" or 平台 == "kona":
        频率 = shell(序列号, "cat /sys/devices/system/cpu/cpu7/cpufreq/cpuinfo_max_freq")
        if 频率.isdigit():
            名称 = "骁龙870" if int(频率) >= 3150000 else ("骁龙865+" if int(频率) >= 3000000 else "骁龙865")
            return f"{名称} ({型号 or 平台})"
    # 骁龙680/685 共用 SM6225，685 大核 2.8GHz、680 为 2.4GHz
    if 型号.upper() == "SM6225":
        频率 = shell(序列号, "cat /sys/devices/system/cpu/cpu7/cpufreq/cpuinfo_max_freq")
        if 频率.isdigit():
            return f"{'骁龙685' if int(频率) >= 2700000 else '骁龙680'} ({型号})"
    # 骁龙888/888+ 共用 SM8350，888+ 大核 2.995GHz、888 为 2.84GHz
    if 型号.upper() == "SM8350":
        频率 = shell(序列号, "cat /sys/devices/system/cpu/cpu7/cpufreq/cpuinfo_max_freq")
        if 频率.isdigit():
            return f"{'骁龙888+' if int(频率) >= 2950000 else '骁龙888'} ({型号})"
    for 代号 in (型号, 平台):
        if 代号 and 代号 in 芯片名称表:
            名称 = 芯片名称表[代号]
            return 名称 if 名称 == 代号 else f"{名称} ({代号})"  # 代号本身就是市场名时不重复
        if 代号 and 代号.upper() in 芯片名称表:
            return f"{芯片名称表[代号.upper()]} ({代号})"
    原始 = " ".join(x for x in (厂商, 型号) if x) or 平台
    return f"{原始}（待查）" if 原始 else ""


def 采集GPU(序列号):
    输出 = shell(序列号, "dumpsys SurfaceFlinger | grep -i GLES")
    匹配 = re.search(r"GLES:\s*(.*?),\s*OpenGL ES", 输出)
    if not 匹配:
        return ""
    文本 = 匹配.group(1)
    # 三星等经 ANGLE 转译：ANGLE ((Samsung Xclipse 540) on Vulkan ...)
    angle = re.search(r"ANGLE \(\((.+?)\)", 文本)
    # 否则形如「厂商, 渲染器」，厂商名本身可能带逗号，取最后一段
    名称 = angle.group(1) if angle else 文本.split(", ")[-1]
    return re.sub(r"\s*\(TM\)\s*", " ", 名称).strip()


def 采集内存(序列号):
    匹配 = re.search(r"MemTotal:\s*(\d+)\s*kB", shell(序列号, "cat /proc/meminfo"))
    if not 匹配:
        return ""
    return 格式化容量(向上取档(int(匹配.group(1)) / 1024 / 1024, 内存档位))


def 采集存储(序列号):
    # 部分机型 df /data 返回的是其他挂载点，改为在完整列表中精确匹配挂载点 /data
    行 = [l for l in shell(序列号, "df -k").splitlines() if l.split() and l.split()[-1] == "/data"]
    try:
        GB = int(行[-1].split()[1]) / 1024 / 1024
    except (IndexError, ValueError):
        return ""
    return 格式化容量(向上取档(GB, 存储档位))


def 采集屏幕(序列号):
    """返回 (分辨率, 密度dpi, 最高刷新率, 最短边dp)"""
    匹配 = re.search(r"Physical size:\s*(\d+)x(\d+)", shell(序列号, "wm size"))
    宽, 高 = (int(匹配.group(1)), int(匹配.group(2))) if 匹配 else (0, 0)
    分辨率 = f"{min(宽, 高)}x{max(宽, 高)}" if 宽 else ""

    匹配 = re.search(r"Physical density:\s*(\d+)", shell(序列号, "wm density"))
    密度 = int(匹配.group(1)) if 匹配 else 0

    帧率 = [float(x) for x in re.findall(r"fps=([\d.]+)", shell(序列号, "dumpsys display | grep -m3 fps="))]
    刷新率 = round(max(帧率)) if 帧率 else 0

    最短边dp = min(宽, 高) * 160 / 密度 if 宽 and 密度 else 0
    return 分辨率, 密度, 刷新率, 最短边dp


def 采集设备类型(序列号, 最短边dp):
    if "tablet" in 取属性(序列号, "ro.build.characteristics"):
        return "平板"
    return "平板" if 最短边dp >= 600 else "手机"


def 采集厂商系统(序列号):
    """返回如「HyperOS OS1.0」「ColorOS V14.0」，原生系统返回空"""
    规则 = [
        ("HyperOS", "ro.mi.os.version.name"),
        ("MIUI", "ro.miui.ui.version.name"),
        ("MagicOS", "ro.build.version.magic"),
        ("HarmonyOS", "hw_sc.build.platform.version"),
        ("EMUI", "ro.build.version.emui"),
        ("ColorOS", "ro.build.version.oplusrom"),
        ("OriginOS", "ro.vivo.os.build.display.id"),
        ("Funtouch/OriginOS", "ro.vivo.os.version"),
    ]
    if not 取属性(序列号, "ro.mi.os.version.name"):
        miui = 取属性(序列号, "ro.miui.ui.version.code")  # 如 14；ui.version.name 是「V140」这种写法
        if miui.isdigit():
            return f"MIUI {miui}"
    for 名称, 属性 in 规则:
        值 = 取属性(序列号, 属性)
        if 值:
            return 值 if 名称.split("/")[0] in 值 else f"{名称} {值}"
    fireos = 取属性(序列号, "ro.build.mktg.fireos")  # 亚马逊 Fire 平板，如「Fire OS 8.3.3.8」
    if fireos:
        return fireos
    oneui = 取属性(序列号, "ro.build.version.oneui")  # 如 60100 → 6.1
    if oneui.isdigit() and len(oneui) >= 5:
        return f"One UI {int(oneui) // 10000}.{int(oneui) // 100 % 100}"
    return ""


def 按型号判断区域(品牌, 型号代号):
    """按各品牌型号代号的区域规则判断，无法判断返回空"""
    品牌 = 品牌.lower()
    代号 = 型号代号.upper().strip()
    if not 代号:
        return ""
    if 品牌 == "google":
        return "海外"  # Pixel 无国行
    # 仅适用数字开头的新式型号（如 23113RKC6C）；「MI 6」「MI 8」等老机型国行海外同名，无法区分
    if 品牌 in ("xiaomi", "redmi", "poco") and re.fullmatch(r"\d{4,5}[A-Z0-9]{4,6}", 代号):
        return "国行" if 代号.endswith("C") else "海外"
    # 仅适用三星手机；平板（SM-X/T/P）Wi-Fi 版全球都以 0 结尾，无法区分
    if 品牌 == "samsung" and 代号.startswith("SM-") and not 代号.startswith(("SM-X", "SM-T", "SM-P")):
        return "国行" if 代号.endswith("0") else "海外"
    if 品牌 in ("huawei", "honor") and "-" in 代号:
        return "国行" if 代号.split("-")[-1][:2] in ("AL", "TL", "AN", "CL", "DL", "UL") else "海外"
    if 品牌 in ("oppo", "oneplus"):
        if 代号.startswith("CPH"):
            return "海外"
        if re.fullmatch(r"P[A-Z]{2,3}\d{2,3}", 代号):
            return "国行"
    if 品牌 in ("vivo", "iqoo") and 代号.startswith("V"):
        return "国行" if 代号.endswith("A") else "海外"
    return ""


def 判断版本区域(序列号, 有GMS, 品牌="", 型号代号=""):
    """返回 (区域, 是否需核对)。优先级：厂商区域属性 → 型号代号规则 → 出厂语言/谷歌服务推测"""
    区域值 = 取属性(序列号, "ro.miui.region", "ro.oplus.regionmark", "ro.vendor.oplus.regionmark",
                 "ro.csc.countryiso_code", "ro.product.locale.region").upper()
    if 区域值:
        return ("国行" if 区域值 in ("CN", "CHN", "CHINA", "PRC") else "海外"), False
    if 取属性(序列号, "ro.vivo.product.overseas").lower() == "yes":
        return "海外", False
    if "_global" in 取属性(序列号, "ro.product.mod_device"):
        return "海外", False
    按型号 = 按型号判断区域(品牌, 型号代号)
    if 按型号:
        return 按型号, False
    出厂语言 = 取属性(序列号, "ro.product.locale")  # 出厂默认语言，非当前设置语言
    if 出厂语言:
        国行 = 出厂语言.startswith("zh") and ("CN" in 出厂语言.upper() or "HANS" in 出厂语言.upper())
    else:
        国行 = not 有GMS
    return ("国行" if 国行 else "海外"), True


def 判断性能档位(CPU):
    名称 = CPU.split(" (")[0]
    for 档位, 芯片列表 in 芯片档位表.items():
        if 名称 in 芯片列表:
            return 档位
    return ""


def 判断屏幕形态(序列号, 设备类型, 屏宽):
    if 设备类型 == "平板":
        return "平板"
    if "hinge_angle" in shell(序列号, "pm list features"):
        return "折叠屏"
    输出 = shell(序列号, "dumpsys display | grep -m1 -o 'boundingRect={[^}]*}'")
    开孔 = [tuple(map(int, m)) for m in
            re.findall(r"Rect\((\d+), (\d+) - (\d+), (\d+)\)", 输出)]
    开孔 = [(l, t, r, b) for l, t, r, b in 开孔 if r > l and b > t]
    if not 开孔:
        return "全面屏无开孔" if "boundingRect" in 输出 else ""
    最宽 = max(r - l for l, t, r, b in 开孔)
    if 屏宽 and 最宽 / 屏宽 > 0.3:
        return "刘海"
    if 屏宽 and 最宽 / 屏宽 > 0.15:
        return "水滴"
    return "挖孔"


def 判断网络版本(序列号):
    """产品名带 wifi 直接判 WLAN 版；否则看是否声明电话功能"""
    if "wifi" in 取属性(序列号, "ro.product.name").lower():
        return "WLAN版"
    特性 = shell(序列号, "pm list features").splitlines()
    return "蜂窝版" if "feature:android.hardware.telephony" in (l.strip() for l in 特性) else "WLAN版"


def 检测手写笔(序列号):
    """检测到手写笔特性或输入设备返回「支持」，否则留空（是否随机附带需人工标注）"""
    # 只认三星 S Pen 标记；摩托罗拉等会声明通用的 com.motorola.stylus，但硬件不一定支持，不作依据
    特性 = shell(序列号, "pm list features").lower()
    if "spen" in 特性:
        return "支持"
    输入设备 = shell(序列号, "dumpsys input | grep -iE 'Location:'")
    if re.search(r"stylus|e-pen|wacom|\bpen\b", 输入设备, re.I):
        return "支持"
    return ""


def 采集蓝牙MAC(序列号):
    值 = shell(序列号, "settings get secure bluetooth_address")
    if 是有效MAC(值):
        return 值.upper()
    匹配 = re.search(MAC正则, shell(序列号, "dumpsys bluetooth_manager | grep -m1 -i 'address:'"))
    if 匹配 and 是有效MAC(匹配.group(1)):
        return 匹配.group(1).upper()
    return ""


def 采集WLAN_MAC(序列号):
    """返回 (MAC, 是否为出厂地址)"""
    匹配 = re.search(MAC正则, shell(序列号, "dumpsys wifi | grep -m1 -i 'factory_mac_address'"))
    if 匹配 and 是有效MAC(匹配.group(1)):
        return 匹配.group(1).upper(), True
    for 命令 in ("cat /sys/class/net/wlan0/address", "ip link show wlan0"):
        匹配 = re.search(MAC正则, shell(序列号, 命令))
        if 匹配 and 是有效MAC(匹配.group(1)):
            return 匹配.group(1).upper(), False
    return "", False


def 采集WebView(序列号):
    匹配 = re.search(r"Current WebView package \(name, version\):\s*\(([^,]+),\s*([^)]+)\)",
                   shell(序列号, "dumpsys webviewupdate"))
    return 匹配.group(2).strip() if 匹配 else ""


def 有谷歌服务(序列号):
    包列表 = shell(序列号, "pm list packages com.google.android.gms").splitlines()
    return any(l.strip() == "package:com.google.android.gms" for l in 包列表)


def 采集单台(adb序列号):
    print(f"\n▶ 正在采集 {adb序列号} ...")
    # 连接中途断开时所有命令都返回空，写进表会把原有正确数据冲掉，直接跳过
    if not 取属性(adb序列号, "ro.build.version.sdk"):
        print(f"⚠ {adb序列号}：设备无响应（可能已断开或未授权），本次跳过，不写入表格")
        return None

    品牌 = 取属性(adb序列号, "ro.product.brand")
    if 品牌.islower():
        品牌 = 品牌.capitalize()  # google → Google，OPPO/HONOR 等原样保留
    型号 = 取属性(adb序列号, "ro.product.model")
    市场名称 = 取属性(
        adb序列号, "ro.product.marketname", "ro.config.marketing_name",
        "ro.vendor.oplus.market.name", "ro.vivo.market.name", "ro.product.vendor.marketname",
        "ro.config.lgsi.market_name", "ro.odm.lenovo.series", "ro.vendor.product.display",
    )
    if not 市场名称 and 品牌.lower() == "amazon":
        市场名称 = 亚马逊型号表.get(型号, "")  # Fire 系统不存市场名，按型号代号对照
    if not 市场名称 and 品牌.lower() == "samsung":
        市场名称 = shell(adb序列号, "settings get global device_name")  # 三星无市场名属性，取出厂设备名称
    市场名称 = 市场名称 if 市场名称 and 市场名称 != "null" else 型号
    for 中文品牌 in ("联想", "小米", "华为", "荣耀", "三星", "一加", "真我"):
        市场名称 = 市场名称.removeprefix(中文品牌).strip()  # 去掉市场名里自带的中文品牌，避免重复
    设备型号 = 市场名称 if 品牌.lower() in 市场名称.lower() else f"{品牌} {市场名称}".strip()
    # 小米、OPPO 的 ro.serialno 只是 adb 标识，产品序列号（包装盒/设置里显示的）另存：
    # 小米在 ro.ril.oem.psno，OPPO 在 vendor.gsm.phoneserial
    序列号 = 取属性(adb序列号, "ro.ril.oem.psno", "vendor.gsm.phoneserial",
                 "ro.serialno", "ro.boot.serialno") or adb序列号
    # 多数品牌 ro.product.model 即型号代号；Pixel 等返回的是带空格的市场名，改取硬件 SKU
    型号代号 = 型号
    if " " in 型号:
        型号代号 = 取属性(adb序列号, "ro.boot.hardware.sku", "ro.boot.product.hardware.sku") or 型号

    分辨率, 密度, 刷新率, 最短边dp = 采集屏幕(adb序列号)
    厂商系统 = 采集厂商系统(adb序列号)
    安卓版本 = 取属性(adb序列号, "ro.build.version.release")
    WLAN_MAC, 是出厂地址 = 采集WLAN_MAC(adb序列号)
    蓝牙MAC = 采集蓝牙MAC(adb序列号)
    架构 = 取属性(adb序列号, "ro.product.cpu.abilist")
    WebView = 采集WebView(adb序列号)

    设备类型 = 采集设备类型(adb序列号, 最短边dp)
    CPU = 采集CPU(adb序列号)
    GMS = 有谷歌服务(adb序列号)
    区域, 区域需核对 = 判断版本区域(adb序列号, GMS, 品牌, 型号代号)
    屏宽 = int(分辨率.split("x")[0]) if 分辨率 else 0

    备注 = []
    if 区域需核对:
        备注.append("版本区域为推测值，请核对")
    if not WLAN_MAC:
        备注.append("系统mac地址无权限读取，请到「状态信息」查看")
    elif not 是出厂地址:
        备注.append("系统mac地址可能为随机地址，请核对")
    if not 蓝牙MAC:
        备注.append("蓝牙MAC无权限读取，请到「状态信息」查看")

    信息 = {
        "设备型号": 设备型号,
        "设备类型": 设备类型,
        "型号代号": 型号代号,
        "序列号": 序列号,
        "操作系统": f"Android {安卓版本}" + (f"（{厂商系统}）" if 厂商系统 else ""),
        "CPU": CPU,
        "GPU": 采集GPU(adb序列号),
        "运行内存": 采集内存(adb序列号),
        "存储容量": 采集存储(adb序列号),
        "分辨率": 分辨率,
        "系统mac地址": WLAN_MAC,
        "API级别": 取属性(adb序列号, "ro.build.version.sdk"),
        "安全补丁": 取属性(adb序列号, "ro.build.version.security_patch"),
        "CPU架构": "arm64" if "arm64" in 架构 else (f"仅32位({架构})" if 架构 else ""),
        "谷歌服务": "有" if GMS else "无",
        "WebView版本": WebView,
        "刷新率": f"{刷新率}Hz" if 刷新率 else "",
        "屏幕密度": f"{密度}dpi" if 密度 else "",
        "蓝牙MAC": 蓝牙MAC,
        "UDID": adb序列号,  # Android 无 UDID，记 adb 序列号作为连电脑时的设备唯一标识
        "版本区域": 区域,
        "网络版本": 判断网络版本(adb序列号),
        "性能档位": 判断性能档位(CPU),
        "屏幕形态": 判断屏幕形态(adb序列号, 设备类型, 屏宽),
        "手写笔": 检测手写笔(adb序列号),
        "采集备注": "；".join(备注),
    }
    for 键, 值 in 信息.items():
        print(f"   {键}：{值 or '（未获取）'}")
    # 采集中途掉线：系统版本、分辨率等关键字段会是空的，整条作废，避免冲掉表里原有数据
    if not 安卓版本 or not 分辨率 or not 取属性(adb序列号, "ro.build.version.sdk"):
        print(f"⚠ {adb序列号}：采集过程中设备断开，数据不完整，本次不写入表格")
        return None
    return 序列号, 信息


def 列出设备():
    已连接, 异常 = [], []
    for 行 in adb(None, "devices").splitlines()[1:]:
        字段 = 行.split()
        if len(字段) >= 2:
            (已连接 if 字段[1] == "device" else 异常).append((字段[0], 字段[1]))
    return 已连接, 异常


def 新建表格():
    工作簿 = Workbook()
    表 = 工作簿.active
    表.title = "测试机"
    表.append(表头)
    for i, 宽 in enumerate(列宽, start=1):
        表.column_dimensions[表.cell(1, i).column_letter].width = 宽
    for 格 in 表[1]:
        格.font = Font(bold=True, color="FFFFFF")
        格.fill = PatternFill("solid", fgColor="4472C4")
        格.alignment = Alignment(horizontal="center", vertical="center")
    表.freeze_panes = "A2"
    # 新增列表头用浅色区分
    for 格 in 表[1][len(现有列):]:
        格.fill = PatternFill("solid", fgColor="70AD47")
    return 工作簿


def 补列(表, 现有表头):
    """旧表缺少后来新增的列时，插到前一列后面，保持与新表一致的列顺序"""
    补了 = False
    for 列名 in 新增列:
        if 列名 in 现有表头:
            continue
        前一列 = 新增列[新增列.index(列名) - 1]
        位置 = 现有表头.index(前一列) + 2 if 前一列 in 现有表头 else len(现有表头) + 1
        表.insert_cols(位置)
        格 = 表.cell(1, 位置, 列名)
        格.font = Font(bold=True, color="FFFFFF")
        格.fill = PatternFill("solid", fgColor="70AD47")
        格.alignment = Alignment(horizontal="center", vertical="center")
        现有表头.insert(位置 - 1, 列名)
        print(f"ℹ 表格新增列：{列名}")
        补了 = True
    if 补了 and 现有表头 == 表头:
        for i, 宽 in enumerate(列宽, start=1):
            表.column_dimensions[表.cell(1, i).column_letter].width = 宽


def 设置下拉(表, 现有表头):
    """重建下拉选项（每次写入都刷新，选项有调整时旧表也能同步）"""
    表.data_validations.dataValidation = []
    for 列名, 选项 in 下拉选项.items():
        字母 = 表.cell(1, 现有表头.index(列名) + 1).column_letter
        校验 = DataValidation(type="list", formula1=f'"{",".join(选项)}"', allow_blank=True)
        校验.add(f"{字母}2:{字母}1000")
        表.add_data_validation(校验)


def 规范名称(值):
    """设备名称比对用：忽略大小写、空格和开头品牌名，如「Apple iPhone 12」=「iPhone 12」"""
    名称 = re.sub(r"\s+", "", str(值 or "")).lower()
    return re.sub(r"^(apple|samsung|三星|google|xiaomi|小米|huawei|华为|lenovo|联想|oppo|vivo|motorola|amazon)", "", 名称)


def 写入表格(路径, 记录列表):
    工作簿 = load_workbook(路径) if os.path.exists(路径) else 新建表格()
    表 = 工作簿.active
    现有表头 = [str(c.value).strip() if c.value else "" for c in 表[1]]
    补列(表, 现有表头)
    现有表头 = [str(c.value).strip() if c.value else "" for c in 表[1]]
    缺失 = [h for h in 表头 if h not in 现有表头]
    if 缺失:
        print(f"✘ 表头缺少列：{'、'.join(缺失)}，请检查 {路径}")
        sys.exit(1)
    设置下拉(表, 现有表头)

    # 按序列号匹配已有行；手工登记、没有序列号的行再按 WLAN MAC 匹配
    序列号列 = 现有表头.index("序列号") + 1
    MAC列 = 现有表头.index("系统mac地址") + 1
    型号列 = 现有表头.index("型号代号") + 1
    资产列 = 现有表头.index("资产编号") + 1
    名称列 = 现有表头.index("设备型号") + 1
    已有行, MAC行, 型号行, 名称行 = {}, {}, {}, {}
    for r in range(2, 表.max_row + 1):
        值 = str(表.cell(r, 序列号列).value or "").strip()
        if 值:
            已有行[值] = r
        # 定制 Pad 以资产编号作序列号，资产编号也作为序列号匹配键
        资产 = str(表.cell(r, 资产列).value or "").strip()
        if 资产:
            已有行.setdefault(资产, r)
        mac = str(表.cell(r, MAC列).value or "").strip().upper().replace("-", ":")
        if mac:
            MAC行[mac] = r
        型号 = str(表.cell(r, 型号列).value or "").strip()
        if 型号:
            型号行.setdefault(型号, []).append(r)
        名称 = 规范名称(表.cell(r, 名称列).value)
        if 名称:
            名称行.setdefault(名称, []).append(r)

    def 按型号匹配(记录):
        """资产清单导入的行没有序列号和 MAC：同型号代号在表中恰好一行且该行未采集过时才合并"""
        for 候选 in (型号行.get(str(记录.get("型号代号", "")).strip(), []),
                   名称行.get(规范名称(记录.get("设备型号")), [])):  # 清单的型号代号列有时填的是市场名
            # 只要求该行还没采集过（无序列号）；手工登记过 MAC 的行也允许合并
            if len(候选) == 1 and not 表.cell(候选[0], 序列号列).value:
                return 候选[0]
        return None

    新增, 更新 = 0, 0
    for 序列号, 记录 in 记录列表:
        # 手工登记时可能把蓝牙 MAC 当成了 WLAN MAC，所以也用蓝牙 MAC 去匹配
        r = (已有行.get(序列号) or MAC行.get(str(记录.get("系统mac地址", "")).upper())
             or MAC行.get(str(记录.get("蓝牙MAC", "")).upper()) or 按型号匹配(记录))
        if r is None:
            r = 表.max_row + 1
            新增 += 1
        else:
            更新 += 1
        for c, 列名 in enumerate(现有表头, start=1):
            if 列名 not in 记录 or 列名 in 手工列:
                continue
            原值 = str(表.cell(r, c).value or "")
            if 列名 in 建议列 and 原值:
                continue  # 建议列已有值（可能人工改过）不覆盖
            if 列名 == "设备型号" and "定制" in 原值:
                continue  # 定制 Pad 保留资产清单里的内部名称
            if 列名 == "分辨率" and "内屏" in 原值:
                continue  # 折叠屏只能读到当前所用屏幕，保留人工核对的内外屏分辨率
            if 列名 == "CPU" and "待查" in 记录[列名] and 原值:
                continue  # 芯片识别不出来时，不覆盖表里已核对的名称
            if 列名 == "CPU" and 记录[列名].split(" (")[0] in 原值 and len(原值) > len(记录[列名].split(" (")[0]):
                continue  # 保留更精确的定制版名称，如「for Galaxy」「8500-Ultra」
            if 列名 == "CPU" and "/" in 记录[列名].split(" (")[0] and 原值:
                候选 = 记录[列名].split(" (")[0]  # 如「天玑6100+/6300」，共用代号无法区分
                前缀 = re.match(r"\D+", 候选).group() if re.match(r"\D+", 候选) else ""
                if any(原值.split(" (")[0] == (x if x.startswith(前缀) else 前缀 + x) for x in 候选.split("/")):
                    continue  # 表里已是其中一个具体型号（人工核对过），不用模糊名覆盖
            if 记录[列名] != "" or not 原值 or 列名 == "采集备注":
                表.cell(r, c).value = 记录[列名]

    try:
        工作簿.save(路径)
    except PermissionError:
        print(f"\n✘ 保存失败：{路径} 正被占用，请先关闭 Excel 后重试")
        sys.exit(1)
    print(f"\n✔ 已写入 {路径}（新增 {新增} 台，更新 {更新} 台）")


def main():
    解析器 = argparse.ArgumentParser(description="采集已连接的 Android / iOS 设备硬件信息并导出 Excel")
    解析器.add_argument("-o", "--output", default=默认输出, help="输出 Excel 路径")
    参数 = 解析器.parse_args()

    已连接, 异常 = 列出设备()
    for 序列号, 状态 in 异常:
        提示 = {"unauthorized": "未授权，请在手机上点「允许 USB 调试」",
                "offline": "离线，请重新插拔数据线"}.get(状态, 状态)
        print(f"⚠ {序列号}：{提示}")
    print(f"检测到 Android 设备 {len(已连接)} 台")
    记录列表 = [记录 for 序列号, _ in 已连接 if (记录 := 采集单台(序列号))]

    sys.path.insert(0, 脚本目录)
    from ios采集 import 采集iOS设备
    iOS记录 = 采集iOS设备()
    print(f"\n检测到 iOS 设备 {len(iOS记录)} 台")
    记录列表 += iOS记录

    if not 记录列表:
        print("✘ 没有检测到可用设备：Android 请开启 USB 调试；iOS 请解锁并点「信任此电脑」")
        sys.exit(1)
    写入表格(参数.output, 记录列表)
    print("提示：Android 的主屏尺寸、蓝牙、上市日期需按官方参数手工填写；资产、借用等列请手工维护；"
          "版本区域、性能档位、屏幕形态等为初值，可人工修改")


if __name__ == "__main__":
    main()
